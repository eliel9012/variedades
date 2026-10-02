#!/usr/bin/env python3
"""Converte a planilha manual de pesquisas de Senador (data/sources/senado-*.xlsx)
em public/data/polls-senado.json.

Diferente de scripts/sync-tse.mjs e scripts/sync-polls.mjs, este script NÃO roda
automaticamente no build: não existe hoje nenhuma fonte pública/gratuita com
endpoint estável para pesquisas de Senador por estado (GCMais e Gazeta do Povo
publicam como texto/matéria, não como API). Os dados aqui vêm de uma planilha
compilada manualmente a partir de matérias jornalísticas reais e citadas
(aba "Sources" da planilha), com data de compilação registrada.

Para atualizar: substitua o arquivo em data/sources/ por uma planilha mais
recente no mesmo formato (abas README/Coverage/Senador/Sources) e rode:
    python3 scripts/convert-senado-xlsx.py data/sources/<arquivo>.xlsx

Requer `openpyxl` (pip install openpyxl).
"""
import json
import re
import sys
from datetime import date, datetime, timezone
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
DEFAULT_SOURCE = ROOT / "data/sources/senado-2026-10-01.xlsx"
OUTPUT = ROOT / "public/data/polls-senado.json"
PRIMARY_GLOB_DIR = ROOT / "data/sources"


# Matéria original por pesquisa, quando a planilha só cita o compilado da
# GCMais (URL com data de 25/08 no caminho, matéria atualizada até 15/09 e sem
# links para as fontes originais). Só entram aqui URLs verificadas em
# 2026-10-02 que batem instituto + UF + período de campo (e números):
#   SP Datafolha 8-10 Sep: Gazeta do Povo, campo 8 a 10/09, 1.610 entrevistas,
#     SP-04189/2026, Marina 13%, Tebet 13%, do Prado 11%, Derrite 10%.
#   MG Datafolha 8-10 Sep: CartaCapital (divulgação Datafolha de 11/09 para SP,
#     MG, RJ, PE e DF), campo 8 a 10/09, Marília Campos 12%.
# Chave: (UF, instituto, campo) exatamente como na aba "Senador".
SOURCE_URL_OVERRIDES = {
    ("SP", "Datafolha", "8–10 Sep 2026"): (
        "https://www.gazetadopovo.com.br/eleicoes/2026/pesquisa-eleitoral-2026/"
        "datafolha-senado-sao-paulo-setembro-2026/"
    ),
    ("MG", "Datafolha", "8–10 Sep 2026"): (
        "https://www.cartacapital.com.br/politica/"
        "datafolha-saiba-as-intencoes-de-voto-para-o-senado-em-sp-rj-mg-pe-e-df/"
    ),
}

def clean(value):
    if value is None:
        return None
    if isinstance(value, str):
        value = value.strip()
        return value or None
    return value


MONTH_ABBR = {
    "jan": 1, "fev": 2, "feb": 2, "mar": 3, "abr": 4, "apr": 4, "mai": 5, "may": 5,
    "jun": 6, "jul": 7, "ago": 8, "aug": 8, "set": 9, "sep": 9, "out": 10, "oct": 10,
    "nov": 11, "dez": 12, "dec": 12,
}

# Nomes completos em português (com e sem acento) e variantes em inglês que
# não se resolvem pelas 3 primeiras letras sozinhas ("Sept.").
MONTH_FULL = {
    "janeiro": 1, "fevereiro": 2, "março": 3, "marco": 3, "abril": 4, "maio": 5,
    "junho": 6, "julho": 7, "agosto": 8, "setembro": 9, "outubro": 10,
    "novembro": 11, "dezembro": 12, "sept": 9,
}


def month_number(raw):
    """Converte nome/abreviação de mês (PT ou EN, com ou sem ponto final) em
    número. Retorna None se não for um mês reconhecido."""
    token = raw.strip().rstrip(".").lower()
    if token in MONTH_FULL:
        return MONTH_FULL[token]
    return MONTH_ABBR.get(token[:3])


# "8 Sep 2026", "8–10 Sep 2026", "28 Aug 2026–2 Sep 2026", "10 Sept. 2026",
# "8 a 10 de setembro de 2026".
TEXT_DATE_RE = re.compile(
    r"(\d{1,2})(?:\s*(?:[–\-]|a)\s*(\d{1,2}))?\s+(?:de\s+)?"
    r"([A-Za-zÀ-ÿ]{3,})\.?\s*(?:de\s+)?(\d{4})"
)
ISO_DATE_RE = re.compile(r"(\d{4})-(\d{2})-(\d{2})")


def parse_fieldwork_end(fieldwork):
    """Extrai (ano, mes, dia-final) de um texto livre de campo, pra comparar
    qual pesquisa é mais recente quando um mesmo estado tem mais de uma na
    planilha (ver pick_most_recent_group). Usa a ÚLTIMA data do texto (o fim
    do campo): "28 Aug 2026–2 Sep 2026" vira 2 Sep. Aceita formato em inglês
    ("8-10 Sep 2026", "10 Sept. 2026"), em português ("8 a 10 de setembro de
    2026") e ISO ("2026-09-10"). Retorna None se nada casar, nunca inventa
    uma data."""
    if not fieldwork:
        return None
    text = str(fieldwork)
    candidates = []
    for match in TEXT_DATE_RE.finditer(text):
        start_day, end_day, month_raw, year = match.groups()
        month = month_number(month_raw)
        if month is None:
            continue
        candidates.append((match.end(), int(year), month, int(end_day or start_day)))
    for match in ISO_DATE_RE.finditer(text):
        year, month, day = (int(x) for x in match.groups())
        candidates.append((match.end(), year, month, day))
    valid = []
    for position, year, month, day in candidates:
        try:
            date(year, month, day)
        except ValueError:
            continue
        valid.append((position, (year, month, day)))
    if not valid:
        return None
    valid.sort(key=lambda item: item[0])
    return valid[-1][1]


def pick_most_recent_group(uf, groups):
    """Quando a planilha tem mais de uma pesquisa pro mesmo estado (achado real:
    SP tinha Datafolha 8-10 Sep e Quaest 4-7 Sep compiladas por engano num único
    card, com candidatos duplicados), mantém só a mais recente por fieldwork,
    igual à convenção já usada para Presidente/Governador (nunca mistura ondas
    de pesquisas diferentes num card só). Nunca derruba a conversão: se a data
    não puder ser comparada, mantém a primeira pesquisa encontrada na planilha."""
    if len(groups) == 1:
        return groups[0]
    dated = [(parse_fieldwork_end(g["fieldwork"]), g) for g in groups]
    if all(d is not None for d, _ in dated):
        dated.sort(key=lambda item: item[0], reverse=True)
        kept = dated[0][1]
    else:
        print(
            f"aviso: {uf} tem {len(groups)} pesquisas na planilha e não foi possível "
            f"comparar as datas de campo; mantendo a primeira encontrada ({groups[0]['pollster']})."
        )
        kept = groups[0]
    discarded = [g for g in groups if g is not kept]
    rewrite_notes_about_discarded(kept, discarded)
    return kept


def rewrite_notes_about_discarded(kept, discarded):
    """As notas da planilha podem citar a outra pesquisa do mesmo estado (achado
    real, SP: "Second SP poll below (Quaest) shown separately; do not compare
    across pollsters."), o que fica falso depois que ela é descartada do card.
    Troca essas notas por um aviso factual montado com os metadados reais da
    pesquisa descartada (instituto e campo como estão na planilha)."""
    for other in discarded:
        other_pollster = (other.get("pollster") or "").strip()
        if not other_pollster or other_pollster == (kept.get("pollster") or "").strip():
            continue
        fieldwork = other.get("fieldwork")
        label = f"{other_pollster} {fieldwork}" if fieldwork else other_pollster
        replacement = f"Outra pesquisa ({label}) existe na planilha; exibida só a mais recente."
        pattern = re.compile(re.escape(other_pollster), re.IGNORECASE)
        for row in kept["results"]:
            if row.get("notes") and pattern.search(row["notes"]):
                row["notes"] = replacement


def generated_at_for(source_path):
    """Data de geração derivada da própria planilha, pra ser determinística
    (mesma planilha, mesmo JSON): primeiro a data yyyy-mm-dd do nome do arquivo
    (ex.: senado-2026-10-01.xlsx), senão o mtime do arquivo em UTC."""
    match = re.search(r"(\d{4})-(\d{2})-(\d{2})", Path(source_path).name)
    if match:
        try:
            day = date(*(int(x) for x in match.groups()))
            return f"{day.isoformat()}T00:00:00Z"
        except ValueError:
            pass
    mtime = datetime.fromtimestamp(Path(source_path).stat().st_mtime, tz=timezone.utc)
    return mtime.replace(microsecond=0).strftime("%Y-%m-%dT%H:%M:%SZ")


def main():
    source_path = Path(sys.argv[1]) if len(sys.argv) > 1 else DEFAULT_SOURCE
    wb = openpyxl.load_workbook(source_path, data_only=True)

    readme_lines = [row[0] for row in wb["README"].iter_rows(values_only=True) if row[0]]

    sources_ws = wb["Sources"]
    source_rows = list(sources_ws.iter_rows(values_only=True))[1:]
    sources = [
        {"publisher": clean(r[0]), "usedFor": clean(r[1]), "url": clean(r[2])}
        for r in source_rows
        if r[0]
    ]

    senador_ws = wb["Senador"]
    rows = list(senador_ws.iter_rows(values_only=True))
    header = rows[0]
    col = {name: idx for idx, name in enumerate(header)}

    # Agrupa por (UF, instituto, campo), não só por UF: a planilha pode ter mais
    # de uma pesquisa pro mesmo estado (achado real: SP tinha Datafolha e Quaest
    # nas mesmas linhas), e misturá-las num card só duplica candidato (mesma
    # "candidateName" duas vezes), o que quebra a renderização em React (key
    # duplicada) e mostra números que não somam 100% de forma nenhuma pesquisa
    # real sozinha.
    by_group = {}
    uf_group_order = {}
    for r in rows[1:]:
        if not r[col["UF"]]:
            continue
        uf = clean(r[col["UF"]])
        pollster = clean(r[col["Pollster"]])
        fieldwork = clean(r[col["Fieldwork"]])
        group_key = (uf, pollster, fieldwork)
        entry = by_group.setdefault(
            group_key,
            {
                "uf": uf,
                "state": clean(r[col["State"]]),
                "pollster": pollster,
                "fieldwork": fieldwork,
                "sample": clean(r[col["Sample"]]),
                "marginOfError": clean(r[col["Margin of error"]]),
                "tseRegistration": clean(r[col["TSE registration"]]),
                "basis": clean(r[col["Basis"]]),
                "completeness": clean(r[col["Completeness"]]),
                "source": clean(r[col["Source"]]),
                "sourceUrl": SOURCE_URL_OVERRIDES.get(group_key) or clean(r[col["URL"]]),
                "results": [],
            },
        )
        if group_key not in uf_group_order.setdefault(uf, []):
            uf_group_order[uf].append(group_key)
        entry["results"].append(
            {
                "candidateName": clean(r[col["Candidate / option"]]),
                "party": clean(r[col["Party"]]),
                "percentage": round(clean(r[col["Vote intention"]]) * 100, 1)
                if isinstance(clean(r[col["Vote intention"]]), (int, float))
                else None,
                "notes": clean(r[col["Notes"]]),
            }
        )

    # Pesquisas conferidas na matéria original (data/sources/senado-primarias-*.json).
    # Entram como mais um grupo por UF: como são mais recentes que as linhas da
    # planilha que citavam só o resumo da GCMais (cujos números não batiam com a
    # fonte primária), pick_most_recent_group fica com elas.
    state_names = {g["uf"]: g["state"] for g in by_group.values()}
    for primary_path in sorted(PRIMARY_GLOB_DIR.glob("senado-primarias-*.json")):
        primary = json.loads(primary_path.read_text(encoding="utf-8"))
        for poll in primary["polls"]:
            uf = poll["uf"]
            group_key = (uf, poll["pollster"], poll["fieldwork"])
            commissioner = poll.get("commissioner")
            by_group[group_key] = {
                "uf": uf,
                "state": state_names.get(uf),
                "pollster": poll["pollster"],
                "fieldwork": poll["fieldwork"],
                "sample": poll["sample"],
                "marginOfError": poll["marginOfError"],
                "tseRegistration": poll["tseRegistration"],
                "basis": poll["basis"],
                "completeness": poll["completeness"],
                "source": f"{poll['pollster']}" + (f" / {commissioner}" if commissioner else ""),
                "sourceUrl": poll["sourceUrl"],
                "results": [
                    {**row, "notes": poll.get("note") if index == 0 else None}
                    for index, row in enumerate(poll["results"])
                ],
            }
            if group_key not in uf_group_order.setdefault(uf, []):
                uf_group_order[uf].append(group_key)
            if not any(s["url"] == poll["sourceUrl"] for s in sources):
                sources.append({"publisher": poll["sourceUrl"].split("/")[2].removeprefix("www."), "usedFor": f"Senado {uf} ({poll['pollster']}, {poll['fieldwork']})", "url": poll["sourceUrl"]})

    states = [
        pick_most_recent_group(uf, [by_group[key] for key in keys])
        for uf, keys in uf_group_order.items()
    ]
    for state in states:
        state["results"].sort(key=lambda row: row["percentage"] or 0, reverse=True)
        total = sum(row["percentage"] or 0 for row in state["results"])
        state["percentageSum"] = round(total, 1)

    states.sort(key=lambda s: s["state"] or s["uf"])

    output = {
        "generatedAt": generated_at_for(source_path),
        "compiledManually": True,
        "note": (
            "Compilação manual a partir de matérias jornalísticas reais (ver "
            "`sources`), não de um feed automático. Não há 2º turno para "
            "Senado no Brasil: cada eleitor vota em até 2 candidatos e os mais "
            "votados são eleitos em turno único."
        ),
        "readme": readme_lines,
        "sources": sources,
        "states": states,
    }

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    OUTPUT.write_text(json.dumps(output, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    print(f"Gravado {OUTPUT} com {len(states)} UFs.")


if __name__ == "__main__":
    main()
